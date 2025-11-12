# ==========================================================
# pages/02_📄 DFD – Formalização da Demanda.py
# SynapseNext – Secretaria de Administração e Abastecimento (TJSP)
# ==========================================================

import os
import json
from io import BytesIO
from docx import Document
import streamlit as st

# ==========================================================
# 📦 Imports institucionais
# ==========================================================
from utils.agents_bridge import AgentsBridge
from utils.integration_dfd import obter_dfd_da_sessao, status_dfd, salvar_dfd_em_json
from utils.ui_components import aplicar_estilo_global, exibir_cabecalho_padrao

# ==========================================================
# ⚙️ Configuração inicial
# ==========================================================
st.set_page_config(page_title="📄 DFD – Formalização da Demanda", layout="wide", page_icon="📄")
aplicar_estilo_global()

exibir_cabecalho_padrao(
    "📄 Formalização da Demanda (DFD)",
    "Pré-preenchimento automático a partir de insumos + geração IA institucional"
)
st.divider()

# ==========================================================
# 🔍 Carregamento automático (sessão + fallback persistente)
# ==========================================================
st.info(status_dfd())
defaults = obter_dfd_da_sessao()

if defaults:
    st.success("📎 Campos do DFD carregados automaticamente do módulo INSUMOS.")
else:
    st.info("Nenhum insumo ativo encontrado. Você pode preencher manualmente ou enviar um documento na aba **🔧 Insumos**.")

# ==========================================================
# 🎨 Estilo institucional SAAB – botões
# ==========================================================
st.markdown("""
<style>
div.stButton > button:first-child {
    background-color: #003366 !important;
    color: white !important;
    border: none !important;
    border-radius: 8px !important;
    height: 2.8em !important;
    font-weight: 500 !important;
}
div.stButton > button:first-child:hover {
    background-color: #002244 !important;
    color: white !important;
    transition: 0.2s;
}
</style>
""", unsafe_allow_html=True)

# ==========================================================
# 🧾 Formulário DFD
# ==========================================================
st.subheader("1️⃣ Entrada – Formalização da Demanda")

with st.form("form_dfd"):
    unidade = st.text_input("Unidade solicitante", value=defaults.get("unidade_solicitante", ""))
    descricao = st.text_area("Descrição da necessidade", value=defaults.get("descricao", ""), height=100)
    motivacao = st.text_area("Justificativa / motivação da contratação", value=defaults.get("justificativa", ""), height=80)
    prazo = st.text_input("Prazo estimado de execução", value=defaults.get("prazo_execucao", ""))
    estimativa_valor = st.text_input("Estimativa de valor", value=defaults.get("estimativa_valor", ""))
    responsavel = st.text_input("Responsável técnico", value=defaults.get("responsavel_tecnico", ""))

    col1, col2 = st.columns(2)
    with col1:
        gerar_ia = st.form_submit_button("⚙️ Gerar rascunho com IA institucional")
    with col2:
        gerar_manual = st.form_submit_button("💾 Gerar rascunho manual")

st.caption("💡 O botão '⚙️ Gerar rascunho com IA institucional' usa o agente DFD.IA para gerar automaticamente o texto técnico.")

# ==========================================================
# 🤖 Geração IA Institucional
# ==========================================================
if gerar_ia:
    st.info("Executando agente DFD institucional...")
    metadata = {
        "unidade": unidade,
        "descricao": descricao,
        "motivacao": motivacao,
        "prazo": prazo,
        "estimativa_valor": estimativa_valor,
        "responsavel": responsavel
    }
    try:
        bridge = AgentsBridge("DFD")
        resultado = bridge.generate(metadata)
        st.success("✅ Rascunho gerado com sucesso pelo agente DFD.IA!")
        st.json(resultado)
        st.session_state["last_dfd"] = resultado.get("secoes", {})
        salvar_dfd_em_json(st.session_state["last_dfd"], origem="ia_dfd")
    except Exception as e:
        st.error(f"Erro ao gerar rascunho com IA: {e}")

# ==========================================================
# ✍️ Geração Manual
# ==========================================================
if gerar_manual:
    dfd_data = {
        "unidade": unidade,
        "descricao": descricao,
        "motivacao": motivacao,
        "prazo": prazo,
        "estimativa_valor": estimativa_valor,
        "responsavel": responsavel
    }
    st.success("✅ Rascunho de DFD gerado manualmente!")
    st.json(dfd_data)
    st.session_state["last_dfd"] = dfd_data
    salvar_dfd_em_json(dfd_data, origem="manual")

# ==========================================================
# 📤 Exportação do Documento
# ==========================================================
if "last_dfd" in st.session_state and st.session_state["last_dfd"]:
    st.divider()
    st.subheader("📤 Exportação de Documento")
    st.info("Baixe o último DFD gerado em formato Word editável.")

    dfd_data = st.session_state["last_dfd"]
    doc = Document()
    doc.add_heading("Formalização da Demanda (DFD)", level=1)
    for k, v in dfd_data.items():
        p = doc.add_paragraph()
        p.add_run(f"{k}: ").bold = True
        p.add_run(str(v) or "—")

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    st.download_button("💾 Baixar DFD_rascunho.docx", buffer, file_name="DFD_rascunho.docx")

    st.markdown("---")
    if st.button("📦 Exportar DFD (JSON)"):
        try:
            path = salvar_dfd_em_json(dfd_data, origem="exportacao_manual")
            st.success(f"✅ DFD exportado com sucesso para {path}")
        except Exception as e:
            st.error(f"Falha ao exportar DFD: {e}")

st.caption("💡 *Dica:* O botão '⚙️ Gerar rascunho com IA institucional' usa o agente DFD.IA para gerar automaticamente o texto técnico.")
